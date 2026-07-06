"""VERIFY — the second amber box of the workflow diagram: "verifica e salva".

Nothing gets marked ready for publishing on the model's word alone: this stage
independently proves every artifact really exists — on disk with the right
format, and (for Notion) retrievable through the API. A product that fails
verification goes back to BUILD; it can never drift to PUBLISH empty-handed.
"""
import logging
import zipfile
from pathlib import Path

from .. import config
from ..integrations import notion
from ..pipeline import pipeline

log = logging.getLogger("factory.verify")

MAGIC = {
    "pdf":  b"%PDF-",
    "xlsx": b"PK\x03\x04",  # xlsx is a zip container
}


def _check_file(artifact: dict) -> str:
    """Returns '' if the artifact is genuinely valid, else the failure reason."""
    path = Path(artifact.get("path", ""))
    if not path.is_file():
        return f"file does not exist: {path}"
    size = path.stat().st_size
    if size < config.MIN_ARTIFACT_BYTES:
        return f"file too small ({size} bytes): {path}"
    kind = artifact.get("kind", "")
    magic = MAGIC.get(kind)
    if magic:
        with path.open("rb") as f:
            head = f.read(len(magic))
        if head != magic:
            return f"wrong magic bytes for {kind}: {path}"
    if kind == "xlsx":
        if not zipfile.is_zipfile(path):
            return f"xlsx is not a valid zip container: {path}"
        try:
            from openpyxl import load_workbook
            wb = load_workbook(path, read_only=True)
            if not wb.sheetnames:
                return f"xlsx has no sheets: {path}"
            wb.close()
        except Exception as e:
            return f"xlsx cannot be opened: {e}"
    return ""


async def run(product: dict) -> bool:
    build = product.get("build") or {}
    artifacts = build.get("artifacts") or []
    checks = []
    failures = []

    if not artifacts:
        failures.append("no artifacts recorded by BUILD")

    for artifact in artifacts:
        reason = _check_file(artifact)
        checks.append({"path": artifact.get("path"),
                       "ok": not reason, "reason": reason or "ok",
                       "bytes": Path(artifact["path"]).stat().st_size
                       if Path(artifact.get("path", "")).is_file() else 0})
        if reason:
            failures.append(reason)

    page_id = build.get("notion_page_id")
    if page_id:
        resp = await notion.retrieve_page(page_id)
        ok = not resp.get("error") and not resp.get("archived", False)
        checks.append({"notion_page_id": page_id, "ok": ok,
                       "reason": resp.get("error", "ok" if ok else "page archived")})
        if not ok:
            failures.append(f"notion page not retrievable: {resp.get('error')}")

    if failures:
        p = pipeline.get(product["id"])
        attempts = (p.get("build_attempts", 0) if p else 0) + 1
        if p:
            p["build_attempts"] = attempts
            p["verify"] = {"verified": False, "checks": checks}
            pipeline.save()
        if attempts >= config.MAX_BUILD_ATTEMPTS:
            pipeline.fail(product["id"], "BUILD_ERROR",
                          f"verification failed {attempts}x: {'; '.join(failures)[:300]}")
        else:
            pipeline.send_back(product["id"], "BUILD",
                               f"verification failed: {'; '.join(failures)[:300]}")
        return False

    pipeline.advance(product["id"], "verify", {"verified": True, "checks": checks})
    log.info("[%s] all artifacts verified (%d checks)", product["id"], len(checks))
    return True
