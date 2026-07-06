"""Real Excel/Google Sheets workbook generation with openpyxl, written to disk.

Turns the DESIGN-stage JSON spec into an actual .xlsx file:
{"sheets": [{"name": ..., "description": ...,
             "columns": [{"header": ..., "example": ..., "formula": ...}]}]}
"""
import logging
import re
from pathlib import Path

log = logging.getLogger("factory.xlsx")

SAMPLE_ROWS = 5


def _clean_sheet_name(name: str, idx: int) -> str:
    # Excel: max 31 chars, no []:*?/\
    name = re.sub(r"[\[\]:*?/\\]", " ", str(name)).strip() or f"Sheet{idx + 1}"
    return name[:31]


def build_xlsx(path: Path, *, title: str, spec: dict) -> Path:
    """Build the workbook and WRITE IT TO DISK. Raises on failure — the caller
    (BUILD stage) must treat any exception as a failed build."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    sheets = spec.get("sheets") or []
    if not sheets:
        raise ValueError("design spec has no sheets")

    wb = Workbook()
    wb.remove(wb.active)

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="2D3A8C")
    title_font = Font(bold=True, size=14, color="2D3A8C")

    # README sheet first: what the buyer downloaded and how to use it
    readme = wb.create_sheet("README")
    readme["A1"] = title
    readme["A1"].font = title_font
    readme["A3"] = "How to use this template:"
    readme["A3"].font = Font(bold=True)
    row = 4
    for i, sheet in enumerate(sheets):
        readme.cell(row=row, column=1,
                    value=f"• {sheet.get('name', f'Sheet{i+1}')}: {sheet.get('description', '')}"[:500])
        row += 1
    readme.cell(row=row + 1, column=1,
                value="Tip: to use in Google Sheets, upload this file to Drive and open with Sheets.")
    readme.column_dimensions["A"].width = 100

    for i, sheet_spec in enumerate(sheets):
        ws = wb.create_sheet(_clean_sheet_name(sheet_spec.get("name", ""), i))
        columns = sheet_spec.get("columns") or []
        if not columns:
            ws["A1"] = sheet_spec.get("description", "")
            continue
        for c, col in enumerate(columns, start=1):
            cell = ws.cell(row=1, column=c, value=str(col.get("header", f"Column {c}"))[:100])
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")
            width = max(14, min(len(str(col.get("header", ""))) + 4, 40))
            ws.column_dimensions[get_column_letter(c)].width = width
        for r in range(2, 2 + SAMPLE_ROWS):
            for c, col in enumerate(columns, start=1):
                formula = col.get("formula")
                example = col.get("example")
                if formula:
                    f = str(formula)
                    if not f.startswith("="):
                        f = "=" + f
                    ws.cell(row=r, column=c, value=f.replace("{row}", str(r)))
                elif example is not None and r == 2:
                    ws.cell(row=r, column=c, value=str(example)[:200])
        ws.freeze_panes = "A2"

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    log.info("XLSX written: %s (%d bytes)", path, path.stat().st_size)
    return path
